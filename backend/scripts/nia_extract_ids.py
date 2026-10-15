"""Development helper: list NIA v2.1 control IDs and domains from a copy of the NCSA standard.

Prints IDs, domain codes and domain names only — never control statement text — so the output can be
committed. Usage: uv run python scripts/nia_extract_ids.py <path to NCSA PDF> [<audit workbook .xlsx>]
"""

import json
import re
import sys
from collections import OrderedDict

import pypdfium2 as pdfium

DOMAIN = re.compile(r"\d+\.\s+([A-Z][A-Za-z&/,\- ]+?)\s*\[([A-Z]{2})\]")
# "XX n. Statement" — an asterisk before the statement marks the control (kept as a flag, not text).
CONTROL = re.compile(r"(?:^|\s)([A-Z]{2}) (\d{1,3})\.\s+(\*?)(?=[A-Z\"(“])")


def pdf_text(path: str) -> str:
    pdf = pdfium.PdfDocument(path)
    return "\n".join(pdf[i].get_textpage().get_text_range() for i in range(len(pdf)))


def main() -> None:
    text = pdf_text(sys.argv[1])
    flat = " ".join(text.split())
    domains: OrderedDict[str, str] = OrderedDict()
    for m in DOMAIN.finditer(flat):
        domains.setdefault(m.group(2), " ".join(m.group(1).split()))
    refs: OrderedDict[str, bool] = OrderedDict()
    for m in CONTROL.finditer(flat):
        if m.group(1) in domains:
            refs.setdefault(f"{m.group(1)} {int(m.group(2))}", m.group(3) == "*")
    extra: list[str] = []
    if len(sys.argv) > 2:
        import openpyxl

        wb = openpyxl.load_workbook(sys.argv[2], read_only=True, data_only=True)
        for row in wb.worksheets[0].iter_rows(min_row=2, values_only=True):
            value = str(row[3] or "").strip()
            m = re.fullmatch(r"([A-Z]{2})\s*(\d{1,3})", value)
            if m and m.group(1) in domains:
                ref = f"{m.group(1)} {int(m.group(2))}"
                if ref not in refs:
                    extra.append(ref)
    print(
        json.dumps(
            {
                "domains": domains,
                "pdf_refs": list(refs),
                "starred": [r for r, s in refs.items() if s],
                "workbook_only": sorted(set(extra)),
            },
            indent=1,
        )
    )


if __name__ == "__main__":
    main()
