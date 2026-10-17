import io

from openpyxl import Workbook, load_workbook

from assayr.audits.tracker import parse_result, read_tracker
from tests.conftest import create_user, sign_in


def tracker_bytes():
    wb = Workbook()
    ws = wb.active
    ws.append(["CONTROL ACTIVITIES (invented)"])
    ws.append([])
    ws.append(["LOGICAL AND PHYSICAL ACCESS CONTROLS"])
    ws.append(["Control Number", "Controls", "Detailed Test of Controls", "Test Results"])
    ws.append(["Criteria: CC6.1: invented criterion text"])
    ws.append(["CC6.1.1", "Invented MFA control", "Inspected invented settings.", "No exceptions noted."])
    ws.append([None, None, "Inspected an invented user list.", None])
    ws.append(
        [" cc 6.1.2 ", "Invented review", "Inspected invented tickets.", "Exception noted: 1 of 10 late"]
    )
    ws.append(["CC6.1.1", "Duplicate", "x", None])
    ws.append(["CC10.1.1", "Unknown criterion", "x", None])
    ws.append(["Note for the auditor", None, None, None])
    ws.append(['=HYPERLINK("x")', None, None, None])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def login(client, db):
    create_user(db, "ana@acme.com")
    sign_in(client, "ana@acme.com")


def upload(client, data, name="Tracker.xlsx"):
    return client.post(
        "/api/v1/audits/import",
        data={"framework": "soc2"},
        files={"file": (name, data, "application/octet-stream")},
    )


def test_parse_result():
    assert parse_result("No exceptions noted.") == ("no_exception", None)
    assert parse_result("no exception noted") == ("no_exception", None)
    assert parse_result("Exceptions noted: 2 samples") == ("exception", "Exceptions noted: 2 samples")
    assert parse_result("Not tested") == ("not_tested", None)
    assert parse_result("  ") == (None, None) and parse_result(None) == (None, None)


def test_read_tracker():
    rows, skipped = read_tracker(tracker_bytes())
    assert [r.code for r in rows] == ["CC6.1.1", "CC6.1.2", "CC10.1.1"]
    assert rows[0].test_procedure == "Inspected invented settings.\nInspected an invented user list."
    assert (rows[0].result, rows[1].result) == ("no_exception", "exception")
    assert rows[1].note == "Exception noted: 1 of 10 late"
    assert any("CC6.1.1" in s and "twice" in s for s in skipped)
    assert any("Note for the auditor" in s for s in skipped)


def test_import_and_export_round_trip(client, db, seeded_all):
    login(client, db)
    r = upload(client, tracker_bytes())
    assert r.status_code == 201, r.text
    body = r.json()
    assert body["imported"] == 2 and body["unknown_refs"] == ["CC10.1.1"]
    detail = client.get("/api/v1/controls/CC6.1.2").json()
    assert detail["test_procedure"] == "Inspected invented tickets." and detail["title"] == "Invented review"
    audit = body["audit"]
    assert audit["period_start"] is None and audit["status"] == "in_progress"
    items = client.get(f"/api/v1/audits/{audit['id']}").json()["items"]
    assert [(i["ref"], i["result"]) for i in items] == [("CC6.1.1", "no_exception"), ("CC6.1.2", "exception")]
    assert items[1]["exception_note"] == "Exception noted: 1 of 10 late"
    assert client.post(f"/api/v1/audits/{audit['id']}/complete").status_code == 409  # no period yet
    client.patch(
        f"/api/v1/audits/{audit['id']}", json={"period_start": "2026-01-01", "period_end": "2026-06-30"}
    )
    assert client.post(f"/api/v1/audits/{audit['id']}/complete").status_code == 200
    xlsx = client.get(f"/api/v1/audits/{audit['id']}/export.xlsx").content
    ws = load_workbook(io.BytesIO(xlsx)).active
    values = [list(row) for row in ws.iter_rows(values_only=True)]
    assert ["Control Number", "Controls", "Detailed Test of Controls", "Test Results"] in [
        v[:4] for v in values
    ]
    assert any(v[0] == "CC6.1.1" and v[3] == "No exceptions noted." for v in values)
    assert any(v[0] == "CC6.1.2" and v[3] == "Exception noted: 1 of 10 late" for v in values)
    assert any(str(v[0] or "").startswith("Criteria: CC6.1:") for v in values)
    assert not any("invented criterion text" in str(c) for v in values for c in v)
    # importing again updates the activities instead of duplicating them
    again = upload(client, tracker_bytes()).json()
    assert again["imported"] == 2
    codes = [c["code"] for c in client.get("/api/v1/controls?framework=soc2&limit=200").json()["items"]]
    assert codes.count("CC6.1.1") == 1


def test_import_with_no_soc2_numbers(client, db, seeded_all):
    login(client, db)
    wb = Workbook()
    wb.active.append(["Control Number", "Controls", "Detailed Test of Controls", "Test Results"])
    wb.active.append(["A.5.1", "x", "y", None])
    buf = io.BytesIO()
    wb.save(buf)
    r = upload(client, buf.getvalue(), "t.xlsx")
    assert r.status_code == 422 and "SOC 2 control numbers" in r.json()["detail"]
    assert upload(client, b"not a workbook").status_code == 422
