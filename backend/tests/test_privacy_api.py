import io
import uuid

from openpyxl import load_workbook
from sqlalchemy import select

from assayr.audit.models import AuditLog
from tests.conftest import create_user, sign_in

ACTIVITY = {
    "name": "Payroll",
    "purpose": "Pay salaries and meet tax duties",
    "data_subjects": ["employees"],
    "data_categories": ["hr", "financial", "national_id"],
    "lawful_basis": "contract",
    "recipients": "Bank",
    "cross_border": False,
    "retention": "7 years after leaving",
    "security_measures": "Role-based access, encryption at rest",
}


def manager(client, db):
    create_user(db, "sana@acme.com", full_name="Sana Malik")
    sign_in(client, "sana@acme.com")


def asset(client, name="HR system", **extra):
    return client.post("/api/v1/assets", json={"name": name, "type": "application", **extra}).json()


def activity(client, **extra):
    response = client.post("/api/v1/privacy/activities", json={**ACTIVITY, **extra})
    assert response.status_code == 201, response.text
    return response.json()


def test_activities_get_sequential_codes(client, db, seeded):
    manager(client, db)
    a = activity(client)
    b = activity(client, name="Recruitment", lawful_basis="consent", data_subjects=["public"])
    assert (a["code"], b["code"]) == ("PA-001", "PA-002")
    assert a["status"] == "draft" and a["sensitive"] is False
    assert [r["code"] for r in client.get("/api/v1/privacy/activities").json()] == ["PA-001", "PA-002"]
    assert [r["code"] for r in client.get("/api/v1/privacy/activities?q=recruit").json()] == ["PA-002"]


def test_unknown_values_are_refused(client, db, seeded):
    manager(client, db)
    for bad in ({"lawful_basis": "because"}, {"data_subjects": ["aliens"]}, {"data_categories": ["dna"]}):
        response = client.post("/api/v1/privacy/activities", json={**ACTIVITY, **bad})
        assert response.status_code == 422, bad


def test_sensitive_and_review_overdue_are_derived(client, db, seeded, clock):
    manager(client, db)
    a = activity(client, data_categories=["health"], status="active", review_due=str(clock.now.date()))
    assert a["sensitive"] is True and a["review_overdue"] is False
    clock.advance(days=1)
    sign_in(client, "sana@acme.com")
    assert client.get(f"/api/v1/privacy/activities/{a['id']}").json()["review_overdue"] is True


def test_linking_assets_marks_them_as_holding_personal_data(client, db, seeded):
    manager(client, db)
    hr = asset(client)
    existing = asset(client, name="Mail", holds_pii=True, pii_categories=["contact"])
    a = activity(client)
    body = client.put(
        f"/api/v1/privacy/activities/{a['id']}/assets", json={"asset_ids": [hr["id"], existing["id"]]}
    )
    assert body.status_code == 200, body.text
    assert sorted(x["name"] for x in body.json()["assets"]) == ["HR system", "Mail"]
    hr_now = next(x for x in client.get("/api/v1/assets").json() if x["name"] == "HR system")
    assert hr_now["holds_pii"] is True and set(hr_now["pii_categories"]) == {"hr", "financial", "national_id"}
    mail_now = next(x for x in client.get("/api/v1/assets").json() if x["name"] == "Mail")
    assert set(mail_now["pii_categories"]) == {"contact", "hr", "financial", "national_id"}
    assert db.scalar(select(AuditLog).where(AuditLog.action == "privacy.assets")) is not None


def test_unknown_assets_are_refused(client, db, seeded):
    manager(client, db)
    a = activity(client)
    response = client.put(
        f"/api/v1/privacy/activities/{a['id']}/assets", json={"asset_ids": [str(uuid.uuid4())]}
    )
    assert response.status_code == 422
    assert client.get(f"/api/v1/privacy/activities/{a['id']}").json()["assets"] == []


def test_inventory_and_summary(client, db, seeded):
    manager(client, db)
    hr = asset(client, confidentiality=5)
    asset(client, name="Clinic records", holds_pii=True, pii_categories=["health"], pii_volume="1k_10k")
    asset(client, name="Web server", type="server")
    a = activity(client, status="active", cross_border=True, transfer_countries="UAE")
    client.put(f"/api/v1/privacy/activities/{a['id']}/assets", json={"asset_ids": [hr["id"]]})
    rows = client.get("/api/v1/privacy/inventory").json()
    assert [r["name"] for r in rows] == ["HR system", "Clinic records"]
    hr_row = rows[0]
    assert hr_row["criticality_label"] == "Very high" and hr_row["activities"] == [
        {"code": "PA-001", "name": "Payroll"}
    ]
    assert rows[1]["pii_sensitive"] is True and rows[1]["pii_volume"] == "1k_10k"
    summary = client.get("/api/v1/privacy/summary").json()
    assert {
        k: summary[k]
        for k in (
            "pii_assets",
            "sensitive_assets",
            "activities",
            "active_activities",
            "cross_border",
            "reviews_due",
        )
    } == {
        "pii_assets": 2,
        "sensitive_assets": 1,
        "activities": 1,
        "active_activities": 1,
        "cross_border": 1,
        "reviews_due": 0,
    }


def test_export_has_both_sheets_and_keeps_text_safe(client, db, seeded):
    manager(client, db)
    activity(client, name="=HYPERLINK(1)", recipients="+cmd")
    response = client.get("/api/v1/privacy/export.xlsx")
    assert response.status_code == 200
    wb = load_workbook(io.BytesIO(response.content))
    assert wb.sheetnames == ["Processing activities", "PII inventory"]
    rows = list(wb["Processing activities"].iter_rows(values_only=True))
    assert rows[0][:4] == ("Code", "Name", "Purpose", "Data subjects")
    assert rows[1][1] == "'=HYPERLINK(1)"


def test_viewers_read_but_do_not_manage(client, db, seeded):
    manager(client, db)
    a = activity(client)
    create_user(db, "vera@acme.com", roles=(("viewer", None),), full_name="Vera")
    sign_in(client, "vera@acme.com")
    assert client.get("/api/v1/privacy/activities").status_code == 200
    assert client.get("/api/v1/privacy/inventory").status_code == 200
    assert client.post("/api/v1/privacy/activities", json=ACTIVITY).status_code == 403
    assert client.patch(f"/api/v1/privacy/activities/{a['id']}", json={"name": "x"}).status_code == 403


def test_activity_edit(client, db, seeded):
    manager(client, db)
    a = activity(client)
    body = client.patch(
        f"/api/v1/privacy/activities/{a['id']}", json={"status": "active", "retention": "10 years"}
    ).json()
    assert body["status"] == "active" and body["retention"] == "10 years"
