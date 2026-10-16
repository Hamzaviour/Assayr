import io

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from assayr.training.models import Assignment
from assayr.training.service import sweep
from assayr.windows.layout import InstallConfig
from assayr.windows.render import render_caddyfile
from tests.conftest import create_user, sign_in


def login(client, email):
    client.cookies.clear()
    client.headers.pop("X-CSRF-Token", None)
    sign_in(client, email)


@pytest.fixture
def setup(client, db, seeded, clock):
    create_user(db, "sana@acme.com", full_name="Sana Malik")
    omar = create_user(db, "omar@acme.com", roles=(("employee", None),), full_name="Omar Ali")
    omar.department = "Finance"
    db.commit()
    login(client, "sana@acme.com")
    c = client.post("/api/v1/courses", json={"title": "Awareness", "recurrence_months": 12}).json()
    client.post(
        f"/api/v1/courses/{c['id']}/materials/drive",
        json={"title": "Video", "url": "https://drive.google.com/file/d/1AbC_dEf-123/view"},
    )
    client.post(f"/api/v1/courses/{c['id']}/publish")
    client.post(f"/api/v1/courses/{c['id']}/assign", json={"department": "Finance"})
    return c


def pass_omar(client, db):
    login(client, "omar@acme.com")
    a = client.get("/api/v1/my/assignments").json()[0]
    d = client.get(f"/api/v1/my/assignments/{a['id']}").json()
    for m in d["materials"]:
        client.post(f"/api/v1/my/assignments/{a['id']}/materials/{m['id']}/done")
    return a["id"]


def open_assignments(db):
    return list(db.scalars(select(Assignment).where(Assignment.renewed.is_(False))))


def test_a_passed_course_comes_back_after_twelve_months(client, db, setup, clock):
    pass_omar(client, db)
    assert sweep(db, clock.now).renewed == 0
    clock.advance(days=364)
    assert sweep(db, clock.now).renewed == 0
    clock.advance(days=3)
    result = sweep(db, clock.now)
    db.commit()
    assert result.renewed == 1
    current = [a for a in open_assignments(db) if a.user.email == "omar@acme.com"]
    assert len(current) == 1 and current[0].status == "assigned" and current[0].source == "renewal"
    assert sweep(db, clock.now).renewed == 0  # idempotent


def test_one_off_and_archived_courses_do_not_come_back(client, db, setup, clock):
    pass_omar(client, db)
    login(client, "sana@acme.com")
    client.post(f"/api/v1/courses/{setup['id']}/archive")
    clock.advance(days=400)
    assert sweep(db, clock.now).renewed == 0


def test_new_joiners_get_the_course_on_the_next_sweep(client, db, setup, clock):
    late = create_user(db, "late@acme.com", roles=(("employee", None),), full_name="Late Joiner")
    late.department = "finance"
    db.commit()
    result = sweep(db, clock.now)
    db.commit()
    assert result.assigned == 1
    assert sweep(db, clock.now).assigned == 0


def test_summary_report_and_dashboard(client, db, setup, clock):
    create_user(db, "nadia@acme.com", roles=(("employee", None),), full_name="Nadia")
    login(client, "sana@acme.com")
    client.post(f"/api/v1/courses/{setup['id']}/assign", json={"target": "all"})
    pass_omar(client, db)
    clock.advance(days=31)
    login(client, "sana@acme.com")
    summary = client.get("/api/v1/training/summary").json()
    row = summary["courses"][0]
    assert row["title"] == "Awareness" and row["assigned"] == 3 and row["passed"] == 1 and row["overdue"] == 2
    assert row["pass_rate"] == pytest.approx(33.3, abs=0.1)
    finance = next(d for d in summary["departments"] if d["department"] == "Finance")
    assert finance == {"department": "Finance", "assigned": 1, "passed": 1, "completion": 100.0}
    report = client.get("/api/v1/training/report.xlsx")
    assert report.status_code == 200
    ws = load_workbook(io.BytesIO(report.content)).active
    rows = list(ws.iter_rows(values_only=True))
    assert ws.title == "Training completion"
    assert list(rows[0]) == [
        "Name",
        "Email",
        "Department",
        "Course",
        "Assigned",
        "Due",
        "Status",
        "Attempts",
        "Score",
        "Completed",
    ]
    assert {r[1] for r in rows[1:]} == {"omar@acme.com", "nadia@acme.com", "sana@acme.com"}
    card = client.get("/api/v1/dashboard").json()["training"]
    assert card == {"assigned": 3, "passed": 1, "overdue": 2}
    login(client, "omar@acme.com")
    assert client.get("/api/v1/training/summary").status_code == 403


def test_dashboard_has_no_training_card_without_the_permission(client, db, seeded):
    create_user(db, "vera@acme.com", roles=(("viewer", None),), full_name="Vera")
    sign_in(client, "vera@acme.com")
    assert client.get("/api/v1/dashboard").json()["training"] is None


def test_the_security_policy_allows_google_drive_frames_only(tmp_path):
    site = (__import__("pathlib").Path(__file__).parents[2] / "frontend" / "site.caddy").read_text(
        encoding="utf-8"
    )
    assert "frame-src https://drive.google.com" in site
    cfg = InstallConfig(
        install_dir=tmp_path / "prog", data_dir=tmp_path / "data", hostname="grc.acme.local", org_name="Acme"
    )
    assert "frame-src https://drive.google.com" in render_caddyfile(cfg)
