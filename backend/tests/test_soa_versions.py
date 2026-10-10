import io

import openpyxl
from sqlalchemy import select

from assayr.audit.models import AuditLog
from tests.conftest import create_user, sign_in


def manager(client, db):
    create_user(db, "sana@acme.com", full_name="Sana Malik")
    sign_in(client, "sana@acme.com")


def executive(client, db):
    create_user(db, "tariq@acme.com", roles=(("executive", None),), full_name="Tariq Qureshi")


def justify_everything(client):
    codes = [i["code"] for i in client.get("/api/v1/soa").json()["items"]]
    client.patch("/api/v1/soa/items", json={"codes": codes, "sources": ["best_practice"]})


def test_submission_needs_clean_checks(client, db, seeded):
    manager(client, db)
    response = client.post("/api/v1/soa/versions")
    assert response.status_code == 409
    assert response.json() == {
        "detail": "Fix the traceability issues before submitting: 93 applicable controls without a justification."
    }


def test_approved_version_is_frozen(client, db, seeded):
    executive(client, db)
    manager(client, db)
    justify_everything(client)
    created = client.post("/api/v1/soa/versions")
    assert created.status_code == 201
    first = created.json()
    assert (first["version_no"], first["status"], first["applicable"], first["excluded"]) == (
        1,
        "pending",
        93,
        0,
    )
    response = client.post("/api/v1/soa/versions")
    assert response.status_code == 409
    assert response.json() == {"detail": "Version 1 is already waiting for approval."}
    sign_in(client, "tariq@acme.com")
    approved = client.post("/api/v1/soa/versions/1/approve").json()
    assert (approved["status"], approved["decided_by"]) == ("approved", "Tariq Qureshi")
    response = client.post("/api/v1/soa/versions/1/approve")
    assert response.status_code == 409
    assert response.json() == {"detail": "Only a version waiting for approval can be approved."}
    assert client.post("/api/v1/soa/versions/1/return", json={"reason": "x"}).status_code == 409
    sign_in(client, "sana@acme.com")
    client.patch(
        "/api/v1/soa/items",
        json={"codes": ["A.8.30"], "applicable": False, "exclusion_reason": "All development is in-house."},
    )
    frozen = client.get("/api/v1/soa/versions/1").json()
    assert next(e for e in frozen["entries"] if e["code"] == "A.8.30")["applicable"] is True
    second = client.post("/api/v1/soa/versions").json()
    assert (second["version_no"], second["excluded"]) == (2, 1)
    assert [v["version_no"] for v in client.get("/api/v1/soa/versions").json()] == [2, 1]
    actions = [e.action for e in db.scalars(select(AuditLog).where(AuditLog.action.like("soa.%")))]
    assert actions.count("soa.submit") == 2 and actions.count("soa.approve") == 1


def test_preparer_cannot_approve_their_own_version(client, db, seeded):
    create_user(
        db, "sana@acme.com", roles=(("isms_manager", None), ("executive", None)), full_name="Sana Malik"
    )
    sign_in(client, "sana@acme.com")
    justify_everything(client)
    client.post("/api/v1/soa/versions")
    response = client.post("/api/v1/soa/versions/1/approve")
    assert response.status_code == 403
    assert response.json() == {"detail": "You prepared this version, so someone else must approve it."}


def test_executive_returns_a_version_with_a_reason(client, db, seeded):
    executive(client, db)
    manager(client, db)
    justify_everything(client)
    client.post("/api/v1/soa/versions")
    sign_in(client, "tariq@acme.com")
    returned = client.post(
        "/api/v1/soa/versions/1/return", json={"reason": "Explain why A.5.7 applies."}
    ).json()
    assert (returned["status"], returned["decision_note"]) == ("returned", "Explain why A.5.7 applies.")
    sign_in(client, "sana@acme.com")
    assert client.post("/api/v1/soa/versions").json()["version_no"] == 2


def test_only_executives_decide(client, db, seeded):
    manager(client, db)
    justify_everything(client)
    client.post("/api/v1/soa/versions")
    assert client.post("/api/v1/soa/versions/1/approve").status_code == 403
    assert client.get("/api/v1/soa/versions/9").status_code == 404


def test_soa_export_lists_every_control(client, db, seeded):
    manager(client, db)
    justify_everything(client)
    client.patch("/api/v1/soa/items", json={"codes": ["A.5.1"], "justification": '=cmd|" /C calc"!A0'})
    client.post("/api/v1/soa/versions")
    response = client.get("/api/v1/soa/versions/1/export.xlsx")
    assert response.status_code == 200 and response.headers["content-disposition"].startswith("attachment")
    book = openpyxl.load_workbook(io.BytesIO(response.content))
    sheet = book["SoA v1"]
    assert sheet.max_row == 94
    assert [c.value for c in sheet[1]][:4] == ["Control", "Title", "Theme", "Applicable"]
    assert (sheet["A2"].value, sheet["D2"].value) == ("A.5.1", "Yes")
    assert sheet["F2"].value == '\'=cmd|" /C calc"!A0'
    approval = {row[0].value: row[1].value for row in book["Approval"].iter_rows(min_row=2)}
    assert (approval["Version"], approval["Status"], approval["Prepared by"]) == (1, "Pending", "Sana Malik")
    working = client.get("/api/v1/soa/export.xlsx")
    assert openpyxl.load_workbook(io.BytesIO(working.content))["SoA (working)"].max_row == 94
