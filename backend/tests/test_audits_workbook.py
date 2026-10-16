import io

import pytest
from openpyxl import Workbook, load_workbook

from assayr.nia.seed import seed_nia
from tests.conftest import create_user, sign_in

NIA_HEADER = [
    "Sr. No.",
    "Domain Name",
    "Sub Domain Name",
    "NIA Control",
    "Control Description/Assessment Area",
    "Design Assessment (DA) / Operating Effectiveness Assessment (OEA) / Both",
    "Mandatory (M) / Additional (A)",
    "Control status (Pass/Fail/Partial)",
]


@pytest.fixture
def nia(db, seeded):
    seed_nia(db, seeded)
    db.commit()
    return seeded


@pytest.fixture
def auditor(client, db, nia):
    create_user(db, "ali@acme.com", roles=(("internal_auditor", None),), full_name="Ali Auditor")
    sign_in(client, "ali@acme.com")


def xlsx(rows: list[list]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "NIA - IA report"
    for row in rows:
        ws.append(row)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def import_(client, data: bytes, framework="nia-2.1"):
    return client.post(
        "/api/v1/audits/import",
        data={"framework": framework},
        files={
            "file": ("audit.xlsx", data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        },
    )


def test_nia_export_matches_the_workbook_layout(client, auditor):
    audit = client.post("/api/v1/audits", json={"framework": "nia-2.1", "title": "NIA audit"}).json()
    client.post(f"/api/v1/audits/{audit['id']}/items", json={"refs": ["SU 5", "SU 7"]})
    items = client.get(f"/api/v1/audits/{audit['id']}").json()["items"]
    client.patch(
        f"/api/v1/audits/{audit['id']}/items/{items[0]['id']}",
        json={"result": "pass", "test_type": "OEA", "assessment_area": "Invented check one"},
    )
    response = client.get(f"/api/v1/audits/{audit['id']}/export.xlsx")
    assert response.status_code == 200 and "attachment" in response.headers["content-disposition"]
    ws = load_workbook(io.BytesIO(response.content)).active
    rows = list(ws.iter_rows(values_only=True))
    assert ws.title == "NIA - IA report" and list(rows[0]) == NIA_HEADER
    assert rows[1][0] == 1 and rows[1][3] == "SU 5" and rows[1][4] == "Invented check one"
    assert rows[1][5] == "OEA" and rows[1][6] in ("M", "A") and rows[1][7] == "Pass"
    assert rows[2][0] == 2 and rows[2][3] == "SU 7" and rows[2][7] is None


def test_iso_export_uses_iso_headings(client, auditor):
    audit = client.post("/api/v1/audits", json={"framework": "iso27001", "title": "ISO audit"}).json()
    client.post(f"/api/v1/audits/{audit['id']}/items", json={"refs": ["A.5.1"]})
    ws = load_workbook(io.BytesIO(client.get(f"/api/v1/audits/{audit['id']}/export.xlsx").content)).active
    header = next(ws.iter_rows(values_only=True))
    assert ws.title == "ISO - IA report" and header[3] == "ISO Control"


def test_import_round_trip_tolerates_messy_headers_and_refs(client, auditor):
    messy = [
        "Sr.  No.",
        "DOMAIN NAME",
        "Sub Domain\nName",
        " nia control ",
        "Control Description/Assessment Area",
        "Design Assessment (DA) /\nOperating Effectiveness Assessment (OEA) / Both",
        "Mandatory (M) / Additional (A)",
        "Control status\n(Pass/Fail/Partial)",
    ]
    data = xlsx(
        [
            messy,
            [1, "System Usage Security", "Usage", "SU5", "Invented area one", "both", "M", "pass"],
            [2, "System Usage Security", "Usage", " su 7 ", "Invented area two", "OEA", "A", "Partial"],
            [3, "Unknown", "Unknown", "ZZ 1", "Invented", None, "M", "Pass"],
            [4, "System Usage Security", "Usage", "SU 5", "Duplicate row", None, "M", "Fail"],
            [5, "System Usage Security", "Usage", "SU 8", "Invented area three", None, None, "NA"],
        [None, None, None, "This invented note explains how the audit was carried out and is not a control."],
        ]
    )
    response = import_(client, data)
    assert response.status_code == 201, response.text
    body = response.json()
    assert body["imported"] == 3 and body["unknown_refs"] == ["ZZ 1"]
    assert body["audit"]["status"] == "in_progress"
    items = {i["ref"]: i for i in client.get(f"/api/v1/audits/{body['audit']['id']}").json()["items"]}
    assert items["SU 5"]["test_type"] == "Both" and items["SU 5"]["result"] == "pass"
    assert items["SU 5"]["assessment_area"] == "Invented area one" and items["SU 5"]["sub_domain"] == "Usage"
    assert items["SU 7"]["result"] == "partial" and items["SU 8"]["result"] == "na"


def test_a_workbook_without_the_expected_header_is_refused(client, auditor):
    response = import_(client, xlsx([["Name", "Value"], ["a", 1]]))
    assert response.status_code == 422 and "Control status" in response.json()["detail"]
    assert client.get("/api/v1/audits").json() == []


def test_a_file_that_is_not_a_workbook_is_refused(client, auditor):
    assert import_(client, b"not a workbook").status_code == 422


def test_an_oversized_workbook_is_refused(client, auditor):
    assert import_(client, b"x" * (5 * 1024 * 1024 + 1)).status_code == 413


def test_viewers_can_export_but_not_import(client, db, nia):
    create_user(db, "vera@acme.com", roles=(("viewer", None),), full_name="Vera")
    sign_in(client, "vera@acme.com")
    assert import_(client, xlsx([NIA_HEADER])).status_code == 403
