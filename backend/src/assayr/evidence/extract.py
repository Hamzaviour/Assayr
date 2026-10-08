"""Identify uploaded files by content and pull out their text for verification."""

from dataclasses import dataclass, field
from pathlib import Path

PDF = "application/pdf"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
TEXT_TYPES = {"text/plain", "text/csv", "text/markdown", "application/json"}
IMAGE_TYPES = {"image/png", "image/jpeg"}
ALLOWED_TYPES = {PDF, DOCX, XLSX, *TEXT_TYPES, *IMAGE_TYPES}
KIND_LABEL = {
    PDF: "PDF",
    DOCX: "Word document",
    XLSX: "Excel workbook",
    "text/plain": "Text file",
    "text/csv": "CSV file",
    "text/markdown": "Markdown file",
    "application/json": "JSON file",
    "image/png": "PNG image",
    "image/jpeg": "JPEG image",
}
MAX_TEXT_CHARS = 2_000_000
_TEXT_EXTENSIONS = {
    ".txt": "text/plain",
    ".csv": "text/csv",
    ".md": "text/markdown",
    ".json": "application/json",
}
_MARKUP_PREFIXES = (b"<html", b"<!doctype", b"<script", b"<svg", b"<?xml")


@dataclass
class Extracted:
    text: str = ""
    pages: int | None = None
    metadata: dict[str, str] = field(default_factory=dict)
    error: str | None = None


def sniff(head: bytes, filename: str) -> str | None:
    """Decide the type from the file's first bytes; the extension only picks between text and zip flavours."""
    name = filename.lower()
    if head.startswith(b"%PDF-"):
        return PDF
    if head.startswith(b"\x89PNG\r\n\x1a\n"):
        return "image/png"
    if head.startswith(b"\xff\xd8\xff"):
        return "image/jpeg"
    if head.startswith(b"PK\x03\x04"):
        if name.endswith(".docx"):
            return DOCX
        if name.endswith(".xlsx"):
            return XLSX
        return None
    if _looks_like_text(head):
        if head.lstrip().lower().startswith(_MARKUP_PREFIXES):
            return None
        for extension, mime in _TEXT_EXTENSIONS.items():
            if name.endswith(extension):
                return mime
    return None


def _looks_like_text(head: bytes) -> bool:
    if not head or b"\x00" in head:
        return False
    return head.decode("utf-8", errors="replace").count("\ufffd") <= 2


def extract(path: Path, mime: str | None) -> Extracted:
    result = _extract(path, mime)
    # NUL can't be stored in a PostgreSQL text column and never helps a rule.
    result.text = result.text.replace("\x00", "")
    result.metadata = {k: v.replace("\x00", "") for k, v in result.metadata.items()}
    return result


def _extract(path: Path, mime: str | None) -> Extracted:
    try:
        if mime == PDF:
            return _pdf(path)
        if mime == DOCX:
            return _docx(path)
        if mime == XLSX:
            return _xlsx(path)
        if mime in TEXT_TYPES:
            return Extracted(text=path.read_text(encoding="utf-8", errors="replace")[:MAX_TEXT_CHARS])
        if mime in IMAGE_TYPES:
            return Extracted()
        return Extracted(error="This file type isn't accepted as evidence.")
    except Exception as exc:  # noqa: BLE001 - any parser failure means the file is unreadable
        if "password" in str(exc).lower():
            return Extracted(error="The PDF is password-protected.")
        return Extracted(error="The file could not be read.")


def _pdf(path: Path) -> Extracted:
    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(str(path))
    try:
        if len(pdf) == 0:
            return Extracted(error="The file could not be read.")
        parts: list[str] = []
        total = 0
        for index in range(len(pdf)):
            page = pdf[index]
            textpage = page.get_textpage()
            text = textpage.get_text_range()
            textpage.close()
            page.close()
            parts.append(text)
            total += len(text)
            if total > MAX_TEXT_CHARS:
                break
        metadata = {k.lower(): str(v) for k, v in pdf.get_metadata_dict(skip_empty=True).items()}
        return Extracted(text="\n".join(parts)[:MAX_TEXT_CHARS], pages=len(pdf), metadata=metadata)
    finally:
        pdf.close()


def _docx(path: Path) -> Extracted:
    import docx

    document = docx.Document(str(path))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    for section in document.sections:
        parts.extend(p.text for p in section.header.paragraphs)
        parts.extend(p.text for p in section.footer.paragraphs)
    props = document.core_properties
    metadata = {
        key: str(value)
        for key, value in {
            "author": props.author,
            "title": props.title,
            "created": props.created,
            "modified": props.modified,
        }.items()
        if value
    }
    return Extracted(text="\n".join(p for p in parts if p)[:MAX_TEXT_CHARS], metadata=metadata)


def _xlsx(path: Path) -> Extracted:
    import openpyxl

    workbook = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    try:
        parts: list[str] = []
        cells = 0
        for sheet in workbook.worksheets:
            parts.append(f"[{sheet.title}]")
            for row in sheet.iter_rows(values_only=True):
                values = [str(v) for v in row if v is not None]
                if values:
                    parts.append(" | ".join(values))
                    cells += len(values)
                if cells > 200_000:
                    break
        return Extracted(text="\n".join(parts)[:MAX_TEXT_CHARS])
    finally:
        workbook.close()
