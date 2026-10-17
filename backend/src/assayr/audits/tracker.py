"""The SOC 2 compliance tracker layout: read a tracker workbook, write a tested audit back in the same layout.

Layout: section headings, a header row "Control Number | Controls | Detailed Test of Controls | Test Results",
"Criteria: ..." rows, one row per control activity, and continuation rows (no number) that add test steps.
"""

import io
from dataclasses import dataclass

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from assayr.audits.models import Audit
from assayr.audits.workbook import MAX_EMPTY_RUN, MAX_ROWS, WorkbookError
from assayr.catalogue.frameworks import load_soc2
from assayr.catalogue.models import Control
from assayr.catalogue.soc2 import criterion_of, normalise_activity
from assayr.exports import safe_cell

HEADER = ["Control Number", "Controls", "Detailed Test of Controls", "Test Results"]
MAX_TEXT = 5000


@dataclass
class TrackerRow:
    code: str
    description: str
    test_procedure: str
    result: str | None
    note: str | None


def _text(row: tuple, index: int) -> str:
    if index >= len(row) or row[index] is None:
        return ""
    return " ".join(str(row[index]).split())


def parse_result(text: str | None) -> tuple[str | None, str | None]:
    """ "No exceptions noted." -> no exception; any other mention of an exception -> exception with the text."""
    clean = " ".join(str(text or "").split())
    lower = clean.lower()
    if not clean:
        return None, None
    if lower.startswith("no exception"):
        return "no_exception", None
    if "not tested" in lower:
        return "not_tested", None
    if "exception" in lower:
        return "exception", clean[:MAX_TEXT]
    return None, None


def read_tracker(data: bytes) -> tuple[list[TrackerRow], list[str]]:
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # any malformed archive or XML is the user's file, not our error
        raise WorkbookError("This file is not an Excel workbook (.xlsx).") from exc
    try:
        return _read(workbook.worksheets[0].iter_rows(values_only=True))
    except WorkbookError:
        raise
    except Exception as exc:
        raise WorkbookError("This workbook could not be read. Save it again from Excel and retry.") from exc
    finally:
        workbook.close()


def _read(rows) -> tuple[list[TrackerRow], list[str]]:
    out: list[TrackerRow] = []
    skipped: list[str] = []
    seen: set[str] = set()
    in_table = False
    current: TrackerRow | None = None
    empty_run = 0
    for number, row in enumerate(rows, start=1):
        first = _text(row, 0)
        if not any(_text(row, i) for i in range(min(len(row), 4))):
            empty_run += 1
            if empty_run >= MAX_EMPTY_RUN:
                break
            continue
        empty_run = 0
        if first.lower() == "control number":
            in_table, current = True, None
            continue
        if not in_table or first.lower().startswith("criteria"):
            current = None if first else current
            continue
        if not first:
            # A continuation row: more test steps (and maybe the result) for the activity above it.
            if current is not None:
                step = _text(row, 2)
                if step:
                    current.test_procedure = f"{current.test_procedure}\n{step}".strip()[:MAX_TEXT]
                if current.result is None and _text(row, 3):
                    current.result, current.note = parse_result(_text(row, 3))
            continue
        code = normalise_activity(first)
        if code is None:
            if first.isupper() and not any(_text(row, i) for i in range(1, 4)):
                current = None  # a section heading between tables
                continue
            skipped.append(f"Row {number}: '{first[:40]}' is not a SOC 2 control number.")
            current = None
            continue
        if code in seen:
            skipped.append(f"Row {number}: {code} appears twice; kept the first.")
            current = None
            continue
        if len(out) >= MAX_ROWS:
            raise WorkbookError(f"The workbook has more than {MAX_ROWS} control activities.")
        result, note = parse_result(_text(row, 3))
        current = TrackerRow(
            code=code,
            description=_text(row, 1)[:MAX_TEXT],
            test_procedure=_text(row, 2)[:MAX_TEXT],
            result=result,
            note=note,
        )
        seen.add(code)
        out.append(current)
    return out, skipped


RESULT_TEXT = {"no_exception": "No exceptions noted.", "not_tested": "Not tested"}


def result_text(result: str | None, note: str | None) -> str:
    if result == "exception":
        text = (note or "").strip()
        return text if text.lower().startswith("exception") else f"Exception noted: {text}".strip()
    return RESULT_TEXT.get(result or "", "")


def export_tracker(audit: Audit, controls: dict[str, Control]) -> Workbook:
    """Writes the audited activities grouped by series and criterion, with Assayr's own criterion summaries."""
    data = load_soc2()
    summaries = {c["ref"]: c["summary"] for c in data["criteria"]}
    by_criterion: dict[str, list] = {}
    for item in audit.items:
        by_criterion.setdefault(criterion_of(item.ref), []).append(item)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "SOC 2 tracker"
    sheet.append([safe_cell(f"SOC 2 control activities: {audit.title}")])
    sheet["A1"].font = Font(bold=True, size=13)
    period = (
        f"Observation period {audit.period_start:%d %b %Y} to {audit.period_end:%d %b %Y}"
        if audit.period_start and audit.period_end
        else "Observation period not set"
    )
    sheet.append([period])
    for series in data["series"]:
        criteria = [
            c["ref"] for c in data["criteria"] if c["series"] == series["ref"] and c["ref"] in by_criterion
        ]
        if not criteria:
            continue
        sheet.append([])
        sheet.append([series["title"].upper()])
        sheet.cell(sheet.max_row, 1).font = Font(bold=True)
        sheet.append(HEADER)
        for cell in sheet[sheet.max_row]:
            cell.font = Font(bold=True)
        for ref in criteria:
            sheet.append([safe_cell(f"Criteria: {ref}: {summaries[ref]}")])
            sheet.cell(sheet.max_row, 1).font = Font(italic=True)
            for item in by_criterion[ref]:
                control = controls.get(item.ref)
                described = (control.description or control.title) if control else ""
                sheet.append(
                    [
                        item.ref,
                        safe_cell(described),
                        safe_cell(item.assessment_area),
                        safe_cell(result_text(item.result, item.exception_note)),
                    ]
                )
    for column, width in zip("ABCD", (16, 60, 60, 30), strict=True):
        sheet.column_dimensions[column].width = width
    return workbook
