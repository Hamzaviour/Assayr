"""Spreadsheet exports. User text never becomes a live formula."""

import io

from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Font

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_FORMULA_START = ("=", "+", "-", "@", "\t", "\r")


def safe_cell(value):
    if isinstance(value, str) and value.startswith(_FORMULA_START):
        return "'" + value
    return value


def add_sheet(workbook: Workbook, title: str, header: list[str], rows: list[list]) -> None:
    sheet = workbook.create_sheet(title)
    sheet.append(header)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append([safe_cell(value) for value in row])
    sheet.freeze_panes = "A2"


def xlsx_response(workbook: Workbook, filename: str) -> Response:
    if "Sheet" in workbook.sheetnames and len(workbook.sheetnames) > 1:
        del workbook["Sheet"]
    buffer = io.BytesIO()
    workbook.save(buffer)
    return Response(
        buffer.getvalue(),
        media_type=XLSX_MIME,
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"',
            "X-Content-Type-Options": "nosniff",
        },
    )
